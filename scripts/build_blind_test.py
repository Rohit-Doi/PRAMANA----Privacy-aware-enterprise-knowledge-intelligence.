"""Build the blind test corpus from its sources.

    python scripts/build_blind_test.py              # documents of the fictional institute
    python scripts/build_blind_test.py --hotpot 100 # plus a sample of HotpotQA (downloads about 50 MB)

Two parts, both unseen by the system until they are ingested:

* ``eval/blind/source`` — documents of a fictional university, written by an independent author
  (see docs/BLIND_TEST.md). They are converted here into the formats named in
  ``eval/blind/documents.jsonl``: Word, PDF, scanned PDF (image only, to be read by OCR), slides.
* HotpotQA — questions written by people that need two Wikipedia paragraphs to answer. For a
  random sample, every paragraph offered with the question (two that matter, eight distractors)
  becomes one document.

Output: files under data/custom/ and rows in data/custom/manifest.jsonl. Then:

    python scripts/ingest_corpus.py --collections blind hotpot
"""

from __future__ import annotations

import argparse
import html
import io
import json
import random
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
BLIND = ROOT / "eval" / "blind"
CUSTOM = ROOT / "data" / "custom"
SCAN_DPI = 200


# ── A small Markdown reader, enough for policy documents ────────────────────────
def blocks_of(markdown: str) -> list[tuple[str, object]]:
    """("heading", (level, text)) | ("table", rows) | ("list", items) | ("text", paragraph)."""
    out: list[tuple[str, object]] = []
    paragraph: list[str] = []
    lines = markdown.replace("\r\n", "\n").split("\n")
    i = 0

    def flush() -> None:
        if paragraph:
            out.append(("text", " ".join(paragraph)))
            paragraph.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            flush()
            out.append(("heading", (len(heading.group(1)), heading.group(2).strip())))
        elif line.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            out.append(("table", rows))
            continue
        elif re.match(r"^\s*(?:[-*]|\d+[.)])\s+", line):
            flush()
            items = []
            while i < len(lines) and re.match(r"^\s*(?:[-*]|\d+[.)])\s+", lines[i]):
                items.append(lines[i].strip())
                i += 1
            out.append(("list", items))
            continue
        elif not line.strip():
            flush()
        else:
            paragraph.append(line.strip())
        i += 1
    flush()
    return out


def plain(text: str) -> str:
    """Drop inline Markdown markup."""
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    return re.sub(r"(\*\*|__|\*|_|`)(?=\S)(.+?)(?<=\S)\1", r"\2", text)


# ── Converters ──────────────────────────────────────────────────────────────────
def to_docx(markdown: str) -> bytes:
    import docx

    document = docx.Document()
    for kind, value in blocks_of(markdown):
        if kind == "heading":
            level, text = value  # type: ignore[misc]
            document.add_heading(plain(text), level=min(level, 4))
        elif kind == "table":
            rows = value  # type: ignore[assignment]
            table = document.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = plain(cell)
        elif kind == "list":
            for item in value:  # type: ignore[union-attr]
                document.add_paragraph(plain(item))
        else:
            document.add_paragraph(plain(str(value)))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def to_pdf(markdown: str) -> bytes:
    """A typeset PDF with ruled tables, as a word processor would export it."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("cell", fontSize=8.5, leading=10.5)
    flow: list[object] = []
    for kind, value in blocks_of(markdown):
        if kind == "heading":
            level, text = value  # type: ignore[misc]
            flow.append(Paragraph(html.escape(plain(text)), styles[f"Heading{min(level, 3)}"]))
        elif kind == "table":
            rows = [[Paragraph(html.escape(plain(c)), cell_style) for c in row] for row in value]  # type: ignore[union-attr]
            table = Table(rows, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            flow.extend([table, Spacer(1, 8)])
        elif kind == "list":
            flow.extend(Paragraph(html.escape(plain(item)), styles["BodyText"]) for item in value)  # type: ignore[union-attr]
        else:
            flow.append(Paragraph(html.escape(plain(str(value))), styles["BodyText"]))
    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=56, rightMargin=56, topMargin=56, bottomMargin=56).build(flow)
    return buffer.getvalue()


def to_scan(markdown: str) -> bytes:
    """A PDF whose pages are pictures of the text: no text layer, as from a scanner."""
    import pymupdf

    source = pymupdf.open(stream=to_pdf(markdown), filetype="pdf")
    scan = pymupdf.open()
    for page in source:
        pixmap = page.get_pixmap(dpi=SCAN_DPI)
        target = scan.new_page(width=page.rect.width, height=page.rect.height)
        target.insert_image(target.rect, pixmap=pixmap)
    return scan.tobytes()


def to_pptx(markdown: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Pt

    presentation = Presentation()
    slide = None
    for kind, value in blocks_of(markdown):
        if kind == "heading":
            _, text = value  # type: ignore[misc]
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = plain(text)
            slide.placeholders[1].text_frame.clear()
            continue
        if slide is None:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = ""
        frame = slide.placeholders[1].text_frame
        lines = (
            [" | ".join(row) for row in value]  # type: ignore[union-attr]
            if kind == "table"
            else list(value)  # type: ignore[call-overload]
            if kind == "list"
            else [str(value)]
        )
        for line in lines:
            paragraph = frame.add_paragraph() if frame.text else frame.paragraphs[0]
            paragraph.text = plain(re.sub(r"^\s*(?:[-*]|\d+[.)])\s+", "", line))
            paragraph.font.size = Pt(14)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


CONVERT = {"docx": to_docx, "pdf": to_pdf, "scan": to_scan, "pptx": to_pptx}
EXTENSION = {"docx": ".docx", "pdf": ".pdf", "scan": ".pdf", "pptx": ".pptx"}


def build_set(name: str) -> list[dict[str, object]]:
    """Convert one authored document set (eval/<name>) into data/custom/<name>."""
    source_dir = ROOT / "eval" / name
    if not (source_dir / "documents.jsonl").exists():
        print(f"{name}: eval/{name}/documents.jsonl not found, skipped")
        return []
    rows = [
        json.loads(line)
        for line in (source_dir / "documents.jsonl").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    out_dir = CUSTOM / name
    out_dir.mkdir(parents=True, exist_ok=True)
    produced: dict[str, str] = {}
    manifest: list[dict[str, object]] = []
    for row in rows:
        source = source_dir / row["file"]
        target_format = row.get("target_format", "md")
        if target_format in CONVERT:
            filename = source.stem + ("-scan" if target_format == "scan" else "") + EXTENSION[target_format]
            (out_dir / filename).write_bytes(CONVERT[target_format](source.read_text(encoding="utf-8")))
        else:
            filename = source.name
            (out_dir / filename).write_bytes(source.read_bytes())
        produced[row["file"]] = f"custom/{name}/{filename}"
    for row in rows:
        manifest.append(
            {
                "path": produced[row["file"]],
                "title": row["title"],
                "department": row["department"],
                "sensitivity": row["sensitivity"],
                "region": row.get("region", "global"),
                "effective_date": row.get("effective_date"),
                "version": str(row.get("version", "1")),
                "supersedes": produced.get(row["supersedes"]) if row.get("supersedes") else None,
                "source_type": "policy",
                "collection": name,
                "license": "Fictional; written for the PRAMANA blind test",
            }
        )
    print(f"{name}: {len(manifest)} documents written to {out_dir}")
    return manifest


def build_hotpot(count: int, seed: int) -> list[dict[str, object]]:
    import os

    os.environ["HF_HUB_OFFLINE"] = "0"
    import pandas as pd
    from huggingface_hub import hf_hub_download

    path = hf_hub_download("hotpotqa/hotpot_qa", "distractor/validation-00000-of-00001.parquet", repo_type="dataset")
    frame = pd.read_parquet(path)
    indexes = list(range(len(frame)))
    random.Random(seed).shuffle(indexes)
    out_dir = CUSTOM / "hotpot"
    out_dir.mkdir(parents=True, exist_ok=True)
    manifest: list[dict[str, object]] = []
    questions: list[dict[str, object]] = []
    seen_titles: dict[str, str] = {}
    for index in indexes[:count]:
        item = frame.iloc[index]
        titles = list(item["context"]["title"])
        paragraphs = [" ".join(s.strip() for s in sentences) for sentences in item["context"]["sentences"]]
        for title, paragraph in zip(titles, paragraphs, strict=True):
            if title in seen_titles or not paragraph.strip():
                continue
            name = re.sub(r"[^A-Za-z0-9]+", "-", title).strip("-")[:80] + f"-{len(seen_titles)}.md"
            (out_dir / name).write_text(f"# {title}\n\n{paragraph}\n", encoding="utf-8")
            seen_titles[title] = name
            manifest.append(
                {
                    "path": f"custom/hotpot/{name}",
                    "title": title,
                    "department": "company",
                    "sensitivity": "PUBLIC",
                    "region": "global",
                    "source_type": "encyclopedia",
                    "collection": "hotpot",
                    "license": "HotpotQA (CC BY-SA 4.0), text from Wikipedia",
                }
            )
        questions.append(
            {
                "id": f"h{len(questions) + 1:03d}",
                "question": str(item["question"]),
                "answer": str(item["answer"]),
                "type": str(item["type"]),
                "level": str(item["level"]),
                "supporting_titles": sorted(set(item["supporting_facts"]["title"])),
            }
        )
    (BLIND / "hotpot_questions.jsonl").write_text(
        "".join(json.dumps(q, ensure_ascii=False) + "\n" for q in questions), encoding="utf-8"
    )
    print(f"hotpot: {len(questions)} questions, {len(manifest)} paragraphs written to {out_dir}")
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--hotpot", type=int, default=0, help="number of HotpotQA questions to sample (0 = none)")
    parser.add_argument("--seed", type=int, default=20261003)
    parser.add_argument("--sets", nargs="+", default=["blind"], help="authored sets under eval/ to build")
    args = parser.parse_args()
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[union-attr]

    CUSTOM.mkdir(parents=True, exist_ok=True)
    manifest_path = CUSTOM / "manifest.jsonl"
    existing = (
        [json.loads(line) for line in manifest_path.read_text(encoding="utf-8").splitlines() if line.strip()]
        if manifest_path.exists()
        else []
    )
    rebuilt = set(args.sets) | ({"hotpot"} if args.hotpot else set())
    rows = [r for r in existing if r.get("collection") not in rebuilt]
    for name in args.sets:
        rows += build_set(name)
    if args.hotpot:
        rows += build_hotpot(args.hotpot, args.seed)
    manifest_path.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows), encoding="utf-8")
    print(f"manifest: {len(rows)} rows in {manifest_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
