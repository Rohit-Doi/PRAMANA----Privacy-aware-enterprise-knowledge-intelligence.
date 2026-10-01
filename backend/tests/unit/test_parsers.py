"""Document parsers: every supported format must yield headings and body text, or a clear error."""

from __future__ import annotations

import io

import pytest

from pramana.knowledge.parsers import detect_mime, parse_document
from pramana.shared.errors import ValidationFailed


def kinds(parsed) -> list[str]:  # type: ignore[no-untyped-def]
    return [b.kind for b in parsed.blocks]


def body(parsed) -> str:  # type: ignore[no-untyped-def]
    return "\n".join(b.text for b in parsed.blocks)


def test_markdown_headings_tables_and_inline_markup() -> None:
    source = (
        "---\ntitle: Leave Policy\n---\n\n"
        "# Leave Policy\n\n"
        "Staff get **25 days** of [annual leave](https://example.com/leave) each year,\n"
        "soft-wrapped onto a second line.\n\n"
        "## Carry-over\n\n"
        "| Region | Days |\n|---|---|\n| EMEA | 5 |\n| APAC | 3 |\n\n"
        "<!-- internal note -->\n"
        "- First item\n- Second item\n"
    )
    parsed = parse_document(source.encode(), filename="leave.md")
    assert parsed.title == "Leave Policy"
    assert "heading" in kinds(parsed) and "table" in kinds(parsed) and "text" in kinds(parsed)
    text = body(parsed)
    assert "25 days of annual leave each year, soft-wrapped onto a second line." in text
    assert "example.com" not in text and "**" not in text and "internal note" not in text
    assert "EMEA" in text and "First item" in text
    headings = [b for b in parsed.blocks if b.kind == "heading"]
    assert [h.level for h in headings] == [1, 2]


def test_plain_text_is_split_into_paragraphs() -> None:
    parsed = parse_document(b"SECURITY RULES\n\nLock your screen.\n\nReport lost devices.", filename="rules.txt")
    assert not parsed.is_empty
    assert "Lock your screen." in body(parsed) and "Report lost devices." in body(parsed)


def test_html_drops_scripts_and_navigation() -> None:
    html = (
        "<html><head><title>Handbook</title><script>var secret = 1;</script></head><body>"
        "<nav>Home | About</nav><h1>Handbook</h1><p>Work hours are flexible.</p>"
        "<table><tr><th>Day</th><th>Hours</th></tr><tr><td>Mon</td><td>8</td></tr></table>"
        "</body></html>"
    )
    parsed = parse_document(html.encode(), filename="handbook.html")
    text = body(parsed)
    assert parsed.title == "Handbook"
    assert "Work hours are flexible." in text
    assert "secret" not in text
    assert "Mon" in text


def test_docx_round_trip() -> None:
    import docx

    document = docx.Document()
    document.add_heading("Expense Policy", level=1)
    document.add_paragraph("Meals are reimbursed up to 60 per day.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Grade", "Limit"
    table.cell(1, 0).text, table.cell(1, 1).text = "L3", "1,200"
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = parse_document(buffer.getvalue(), filename="expenses.docx")
    assert "heading" in kinds(parsed)
    assert "Meals are reimbursed up to 60 per day." in body(parsed)
    assert "1,200" in body(parsed)


def test_pdf_text_is_extracted_with_page_numbers() -> None:
    import pymupdf

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 90), "Data Retention Standard", fontsize=20)
    page.insert_text((72, 140), "Backups are kept for ninety days.", fontsize=11)
    page.insert_text((72, 160), "Audit logs are kept for seven years.", fontsize=11)
    data = pdf.tobytes()
    pdf.close()

    parsed = parse_document(data, filename="retention.pdf")
    text = body(parsed)
    assert "Backups are kept for ninety days." in text
    assert "Audit logs are kept for seven years." in text
    assert {b.page for b in parsed.blocks} == {1}


def test_pdf_without_a_text_layer_is_rejected_with_a_reason() -> None:
    import pymupdf

    pdf = pymupdf.open()
    pdf.new_page()
    data = pdf.tobytes()
    pdf.close()
    with pytest.raises(ValidationFailed) as error:
        parse_document(data, filename="scan.pdf")
    assert "OCR" in str(error.value.detail if hasattr(error.value, "detail") else error.value)


def test_email_headers_and_body() -> None:
    raw = (
        b"From: Ana Lima <ana@example.com>\r\nTo: ops@example.com\r\nSubject: Outage follow-up\r\n"
        b"Date: Tue, 03 Mar 2026 10:00:00 +0000\r\nContent-Type: text/plain; charset=utf-8\r\n\r\n"
        b"The database failover took four minutes.\r\n"
    )
    parsed = parse_document(raw, filename="outage.eml")
    assert parsed.title == "Email: Outage follow-up"
    assert "The database failover took four minutes." in body(parsed)


def test_unsupported_and_corrupt_files_are_rejected() -> None:
    with pytest.raises(ValidationFailed):
        parse_document(b"MZ\x90\x00", filename="tool.exe")
    with pytest.raises(ValidationFailed):
        parse_document(b"this is not a pdf", filename="broken.pdf")


def test_mime_detection_uses_the_extension() -> None:
    assert detect_mime("a.PDF") == "application/pdf"
    assert detect_mime("notes.md").startswith("text/")
