"""Document parsers. Every format is reduced to the same structure: a title and a list of
blocks (headings with a level, body text, tables), each optionally tied to a page."""

from __future__ import annotations

import email
import email.policy
import io
import re
import statistics
from collections import Counter
from dataclasses import dataclass, field
from pathlib import PurePath

from bs4 import BeautifulSoup

from pramana.shared.errors import ValidationFailed
from pramana.shared.text import normalize_text


@dataclass(slots=True)
class Block:
    kind: str  # "heading" | "text" | "table"
    text: str
    level: int = 0
    page: int | None = None


@dataclass(slots=True)
class ParsedDocument:
    title: str
    blocks: list[Block]
    meta: dict[str, object] = field(default_factory=dict)

    @property
    def is_empty(self) -> bool:
        return not any(b.text.strip() for b in self.blocks if b.kind != "heading")


SUPPORTED_EXTENSIONS = {
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".txt": "text/plain",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".html": "text/html",
    ".htm": "text/html",
    ".eml": "message/rfc822",
    ".csv": "text/csv",
}


def detect_mime(filename: str) -> str:
    extension = PurePath(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValidationFailed(
            f"Unsupported file type '{extension}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    return SUPPORTED_EXTENSIONS[extension]


def parse_document(data: bytes, *, filename: str, mime_type: str | None = None) -> ParsedDocument:
    mime = mime_type or detect_mime(filename)
    fallback_title = re.sub(r"[_\-]+", " ", PurePath(filename).stem).strip().title()
    if mime == "text/markdown":
        parsed = _parse_markdown(_decode(data))
    elif mime == "text/plain":
        parsed = _parse_plain(_decode(data))
    elif mime == "application/pdf":
        parsed = _parse_pdf(data)
    elif mime.endswith("wordprocessingml.document"):
        parsed = _parse_docx(data)
    elif mime.endswith("presentationml.presentation"):
        parsed = _parse_pptx(data)
    elif mime == "text/html":
        parsed = _parse_html(_decode(data))
    elif mime == "message/rfc822":
        parsed = _parse_email(data)
    elif mime == "text/csv":
        parsed = _parse_plain(_decode(data))
    else:
        raise ValidationFailed(f"Unsupported content type: {mime}")
    if not parsed.title.strip():
        parsed.title = fallback_title
    parsed.blocks = [b for b in parsed.blocks if b.text.strip()]
    return parsed


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


# ── Markdown ────────────────────────────────────────────────────────────────
_FRONT_MATTER = re.compile(r"\A---\s*\n(.*?)\n---\s*\n", re.DOTALL)
_FM_TITLE = re.compile(r"^title:\s*[\"']?(.+?)[\"']?\s*$", re.MULTILINE)
_SHORTCODE_BLOCK = re.compile(r"\{\{[<%]\s*/?\s*[\w\-/]+[^}]*[>%]\}\}")
_HTML_COMMENT = re.compile(r"<!--.*?-->", re.DOTALL)
_MD_IMAGE = re.compile(r"!\[[^\]]*\]\([^)]*\)")
_MD_LINK = re.compile(r"\[([^\]]+)\]\((?:[^)(]|\([^)]*\))*\)")
_MD_REF_LINK = re.compile(r"\[([^\]]+)\]\[[^\]]*\]")
_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_MD_EMPHASIS = re.compile(r"(\*\*|__|\*|_|`)(?=\S)(.+?)(?<=\S)\1")
_HTML_TAG = re.compile(r"</?[a-zA-Z][^>]*>")
_KRAMDOWN_ATTR = re.compile(r"\{:\s*[^}]*\}|\{#[\w\-]+\}")
_MD_LIST_ITEM = re.compile(r"(?:- |\d{1,3}[.)] )")


def _clean_inline_markdown(line: str) -> str:
    line = _MD_IMAGE.sub("", line)
    line = _MD_LINK.sub(r"\1", line)
    line = _MD_REF_LINK.sub(r"\1", line)
    line = _KRAMDOWN_ATTR.sub("", line)
    line = _HTML_TAG.sub("", line)
    for _ in range(2):
        line = _MD_EMPHASIS.sub(r"\2", line)
    return line


def _parse_markdown(text: str) -> ParsedDocument:
    title = ""
    front = _FRONT_MATTER.match(text)
    if front:
        found = _FM_TITLE.search(front.group(1))
        if found:
            title = found.group(1).strip()
        text = text[front.end() :]
    text = _HTML_COMMENT.sub("", text)
    text = _SHORTCODE_BLOCK.sub("", text)

    blocks: list[Block] = []
    buffer: list[str] = []
    table: list[str] = []
    in_code = False

    def flush_text() -> None:
        if buffer:
            lines: list[str] = []
            for item in buffer:
                if lines and not _MD_LIST_ITEM.match(item) and not _MD_LIST_ITEM.match(lines[-1]):
                    lines[-1] = f"{lines[-1]} {item}"  # soft-wrapped line of the same paragraph
                else:
                    lines.append(item)
            blocks.append(Block("text", "\n".join(lines)))
            buffer.clear()

    def flush_table() -> None:
        if table:
            rows = [r for r in table if not re.fullmatch(r"\|?[\s:\-|]+\|?", r)]
            blocks.append(Block("table", "\n".join(rows)))
            table.clear()

    for raw in text.split("\n"):
        stripped = raw.strip()
        if stripped.startswith("```") or stripped.startswith("~~~"):
            in_code = not in_code
            continue
        if in_code:
            buffer.append(raw.rstrip())
            continue
        heading = _MD_HEADING.match(stripped)
        if heading:
            flush_text()
            flush_table()
            heading_text = _clean_inline_markdown(heading.group(2)).strip()
            if heading_text:
                level = len(heading.group(1))
                if not title and level == 1:
                    title = heading_text
                blocks.append(Block("heading", heading_text, level=level))
            continue
        if stripped.startswith("|") and stripped.count("|") >= 2:
            flush_text()
            table.append(_clean_inline_markdown(stripped))
            continue
        flush_table()
        if not stripped:
            flush_text()
            continue
        cleaned = _clean_inline_markdown(stripped)
        cleaned = re.sub(r"^[-*+]\s+", "- ", cleaned)
        if cleaned.strip():
            buffer.append(cleaned)
    flush_text()
    flush_table()
    return ParsedDocument(title=title, blocks=blocks)


# ── Plain text (contracts and similar) ──────────────────────────────────────
_PLAIN_HEADING = re.compile(
    r"^(?:(?:ARTICLE|SECTION|EXHIBIT|SCHEDULE|APPENDIX)\s+[\dIVXLA-Z]+\b.*"
    r"|\d{1,2}(?:\.\d{1,2})?\.?\s+[A-Z][A-Za-z ,&/\-]{2,70}"
    r"|[A-Z][A-Z0-9 ,&/\-']{4,70})$"
)


def _parse_plain(text: str) -> ParsedDocument:
    blocks: list[Block] = []
    buffer: list[str] = []
    title = ""

    def flush() -> None:
        if buffer:
            blocks.append(Block("text", " ".join(buffer)))
            buffer.clear()

    for raw in text.split("\n"):
        line = raw.strip()
        if not line:
            flush()
            continue
        words = line.split()
        if len(words) <= 12 and _PLAIN_HEADING.match(line) and not line.endswith((".", ",", ";")):
            flush()
            blocks.append(Block("heading", line.title() if line.isupper() else line, level=2))
            if not title:
                title = line.title() if line.isupper() else line
            continue
        buffer.append(line)
    flush()
    return ParsedDocument(title=title, blocks=blocks)


# ── PDF ─────────────────────────────────────────────────────────────────────
def _parse_pdf(data: bytes) -> ParsedDocument:
    import pymupdf

    try:
        pdf = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ValidationFailed(f"The PDF could not be opened: {exc}") from exc

    # One entry per layout block: page, lines as (font size, bold, text).
    layout: list[tuple[int, list[tuple[float, bool, str]]]] = []
    pages_without_text = 0
    with pdf:
        page_count = pdf.page_count
        title = (pdf.metadata or {}).get("title", "") or ""
        for page_index, page in enumerate(pdf, start=1):
            found_text = False
            # Unsorted extraction follows the file's own reading order, which keeps the
            # columns of multi-column pages apart; sorting by position would interleave them.
            for block in page.get_text("dict").get("blocks", []):
                if block.get("type") != 0:
                    continue
                block_lines: list[tuple[float, bool, str]] = []
                for line in block.get("lines", []):
                    spans = [s for s in line.get("spans", []) if s.get("text", "").strip()]
                    if not spans:
                        continue
                    text = "".join(s["text"] for s in line["spans"]).strip()
                    size = max(s["size"] for s in spans)
                    bold = all("bold" in s.get("font", "").lower() or s.get("flags", 0) & 16 for s in spans)
                    block_lines.append((round(size, 1), bold, text))
                if block_lines:
                    layout.append((page_index, block_lines))
                    found_text = True
            if not found_text:
                pages_without_text += 1

    all_lines = [line for _, block_lines in layout for line in block_lines]
    if not all_lines:
        raise ValidationFailed(
            "This PDF has no text layer (it looks scanned). Run OCR on it first; OCR is not enabled here."
        )

    # Running headers and footers repeat on most pages: drop them.
    shape_counts = Counter(re.sub(r"\d+", "#", text) for _, _, text in all_lines)
    repeated = {
        shape
        for shape, n in shape_counts.items()
        if page_count >= 4 and n >= max(3, page_count * 0.5) and len(shape) < 90
    }
    body_size = statistics.median(size for size, _, _ in all_lines)
    page_number = re.compile(r"(?:page\s*)?\d{1,4}(?:\s*of\s*\d+)?", re.I)

    blocks: list[Block] = []
    for page_index, block_lines in layout:
        kept = [
            (size, bold, text)
            for size, bold, text in block_lines
            if re.sub(r"\d+", "#", text) not in repeated and not page_number.fullmatch(text)
        ]
        if not kept:
            continue
        text = _join_pdf_lines([t for _, _, t in kept])
        size = max(s for s, _, _ in kept)
        bold = all(b for _, b, _ in kept)
        words = len(text.split())
        # A heading is a short block set larger than body text, or a short bold block.
        larger = size >= body_size * 1.12
        is_heading = (
            len(kept) <= 2
            and words <= 14
            and not text.endswith((".", ",", ";", ":", "-"))
            and (larger or (bold and words <= 10))
            and re.search(r"[A-Za-z]{3}", text) is not None
        )
        previous = blocks[-1] if blocks else None
        if is_heading:
            level = 1 if size >= body_size * 1.5 else 2 if size >= body_size * 1.25 else 3
            blocks.append(Block("heading", text, level=level, page=page_index))
            if not title and level == 1:
                title = text
        elif (
            previous is not None
            and previous.kind == "text"
            and text[:1].islower()
            and not previous.text.endswith((".", ":", ";", "?", "!"))
        ):
            # A paragraph continued in the next layout block (line group, column or page).
            previous.text = _join_pdf_lines([previous.text, text])
        else:
            blocks.append(Block("text", text, page=page_index))
    return ParsedDocument(
        title=title.strip(),
        blocks=blocks,
        meta={"pages": page_count, "pages_without_text": pages_without_text},
    )


def _join_pdf_lines(lines: list[str]) -> str:
    """Re-flow hard-wrapped lines, repairing words hyphenated across a line break."""
    out = lines[0]
    for line in lines[1:]:
        if out.endswith("-") and line[:1].islower():
            out = out[:-1] + line
        else:
            out = out + " " + line
    return out


# ── DOCX ────────────────────────────────────────────────────────────────────
def _parse_docx(data: bytes) -> ParsedDocument:
    from docx import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise ValidationFailed(f"The DOCX file could not be opened: {exc}") from exc
    title = (document.core_properties.title or "").strip()
    blocks: list[Block] = []
    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
            if style == "title":
                title = title or text
                blocks.append(Block("heading", text, level=1))
            elif style.startswith("heading"):
                digits = re.findall(r"\d+", style)
                blocks.append(Block("heading", text, level=min(int(digits[0]) + 1, 6) if digits else 2))
            elif "list" in style:
                blocks.append(Block("text", "- " + text))
            else:
                blocks.append(Block("text", text))
        elif tag == "tbl":
            table = Table(element, document)
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            blocks.append(Block("table", "\n".join(f"| {r} |" for r in rows)))
    return ParsedDocument(title=title, blocks=blocks)


# ── PPTX ────────────────────────────────────────────────────────────────────
def _parse_pptx(data: bytes) -> ParsedDocument:
    from pptx import Presentation

    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise ValidationFailed(f"The PPTX file could not be opened: {exc}") from exc
    blocks: list[Block] = []
    title = (presentation.core_properties.title or "").strip()
    for index, slide in enumerate(presentation.slides, start=1):
        slide_title = (
            slide.shapes.title.text.strip()
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame
            else ""
        )
        blocks.append(Block("heading", slide_title or f"Slide {index}", level=2, page=index))
        title = title or slide_title
        for shape in slide.shapes:
            if shape == slide.shapes.title or not shape.has_text_frame:
                continue
            text = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
            if text:
                blocks.append(Block("text", text, page=index))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                blocks.append(Block("text", "Speaker notes: " + notes, page=index))
    return ParsedDocument(title=title, blocks=blocks)


# ── HTML ────────────────────────────────────────────────────────────────────
def _parse_html(text: str) -> ParsedDocument:
    soup = BeautifulSoup(text, "lxml")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "noscript", "form"]):
        tag.decompose()
    title = soup.title.get_text(strip=True) if soup.title else ""
    root = soup.find("main") or soup.find("article") or soup.body or soup
    blocks: list[Block] = []
    for element in root.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table", "pre"]):
        if element.name == "table":
            rows = [
                " | ".join(c.get_text(" ", strip=True) for c in row.find_all(["th", "td"]))
                for row in element.find_all("tr")
            ]
            blocks.append(Block("table", "\n".join(f"| {r} |" for r in rows if r.strip())))
        elif element.find_parent("table") is not None:
            continue
        elif element.name.startswith("h"):
            heading = element.get_text(" ", strip=True)
            if heading:
                blocks.append(Block("heading", heading, level=int(element.name[1])))
                title = title or heading
        else:
            body = element.get_text(" ", strip=True)
            if body:
                blocks.append(Block("text", ("- " if element.name == "li" else "") + body))
    return ParsedDocument(title=title, blocks=blocks)


# ── Email ───────────────────────────────────────────────────────────────────
def _parse_email(data: bytes) -> ParsedDocument:
    message = email.message_from_bytes(data, policy=email.policy.default)
    subject = str(message.get("Subject", "")).strip()
    sender = str(message.get("From", "")).strip()
    recipients = str(message.get("To", "")).strip()
    sent = str(message.get("Date", "")).strip()
    body_part = message.get_body(preferencelist=("plain", "html"))
    body = body_part.get_content() if body_part is not None else ""
    if body_part is not None and body_part.get_content_type() == "text/html":
        body = BeautifulSoup(body, "lxml").get_text("\n")
    header = f"Email from {sender or 'unknown sender'} to {recipients or 'unknown recipients'}"
    if sent:
        header += f", sent {sent}"
    blocks = [Block("text", header + f". Subject: {subject or '(none)'}.")]
    for paragraph in re.split(r"\n\s*\n", normalize_text(body)):
        if paragraph.strip():
            blocks.append(Block("text", " ".join(paragraph.split())))
    return ParsedDocument(
        title=f"Email: {subject}" if subject else "Email",
        blocks=blocks,
        meta={"from": sender, "to": recipients, "date": sent},
    )
